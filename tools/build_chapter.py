#!/usr/bin/env python3
"""Build a .docx and a .md copy of a chapter for sending or reading.

Usage: python3 tools/build_chapter.py chapters/chapter-04.md [outdir]
Needs python-docx (pip install python-docx). Output goes to outdir (default
./build, which is git-ignored by habit; do not commit it).
"""
import os
import re
import shutil
import sys

from docx import Document

TOKEN = re.compile(r"(\*\*\*.+?\*\*\*|\*\*.+?\*\*|\*.+?\*)")


def build(src, outdir):
    os.makedirs(outdir, exist_ok=True)
    doc = Document()
    for line in open(src, encoding="utf-8").read().split("\n"):
        if not line.strip():
            continue
        if line.startswith("# "):
            doc.add_heading(line[2:], level=1)
            continue
        para = doc.add_paragraph()
        for part in TOKEN.split(line):
            if not part:
                continue
            if part.startswith("***") and part.endswith("***") and len(part) > 6:
                run = para.add_run(part[3:-3]); run.bold = run.italic = True
            elif part.startswith("**") and part.endswith("**") and len(part) > 4:
                run = para.add_run(part[2:-2]); run.bold = True
            elif part.startswith("*") and part.endswith("*") and len(part) > 2:
                run = para.add_run(part[1:-1]); run.italic = True
            else:
                para.add_run(part)
    stem = os.path.splitext(os.path.basename(src))[0]
    doc.save(os.path.join(outdir, stem + ".docx"))
    shutil.copy(src, os.path.join(outdir, stem + ".md"))
    print("built", os.path.join(outdir, stem + ".docx"))


if __name__ == "__main__":
    if len(sys.argv) < 2:
        sys.exit(__doc__)
    build(sys.argv[1], sys.argv[2] if len(sys.argv) > 2 else "build")
